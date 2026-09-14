"""Page-text extraction for PDF / DOCX / plain-text documents."""
from __future__ import annotations

import io
import tempfile
import zipfile
from xml.etree import ElementTree

from aaa.tools.client_doc_ingest.xml_safety import reject_dtd


def _extract_pdf_pages(data: bytes) -> list[tuple[int | None, str]]:
    """Extract per-page text from a PDF using pypdfium2."""
    import pypdfium2  # type: ignore

    pages: list[tuple[int | None, str]] = []
    with tempfile.NamedTemporaryFile(suffix=".pdf") as tmp:
        tmp.write(data)
        tmp.flush()
        doc = pypdfium2.PdfDocument(tmp.name)
        try:
            for idx, page in enumerate(doc, start=1):
                textpage = page.get_textpage()
                try:
                    pages.append((idx, textpage.get_text_bounded() or ""))
                finally:
                    textpage.close()
                    page.close()
        finally:
            doc.close()
    return pages


def _extract_docx_pages(data: bytes) -> list[tuple[int | None, str]]:
    """Extract text from a DOCX via python-docx, falling back to raw XML."""
    try:
        import docx  # type: ignore

        document = docx.Document(io.BytesIO(data))
        text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
        return [(None, text)]
    except ImportError:
        pass

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        xml = archive.read("word/document.xml")
    reject_dtd(xml)
    root = ElementTree.fromstring(xml)
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    text = "\n".join(node.text or "" for node in root.iter(f"{namespace}t"))
    return [(None, text)]


def _extract_pages(data: bytes, content_type: str) -> list[tuple[int | None, str]]:
    """Dispatch extraction by content type (default: UTF-8 decode)."""
    if content_type == "pdf":
        return _extract_pdf_pages(data)
    if content_type == "docx":
        return _extract_docx_pages(data)
    return [(None, data.decode("utf-8", errors="replace"))]
